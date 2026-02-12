"""
Script to generate GL Upload Template XLSX file
This creates a template file with the expected headers for GL transaction uploads
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

def generate_gl_template():
    """Generate a GL upload template with proper headers and formatting"""
    
    wb = Workbook()
    ws = wb.active
    ws.title = "GL Transactions"
    
    # Define the headers based on XLSX_TO_GL_TRANSACTION_MAP
    # Row 1: Title
    ws['A1'] = "General Ledger Upload Template"
    ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
    ws['A1'].fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    ws['A1'].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells('A1:Y1')
    ws.row_dimensions[1].height = 25
    
    # Row 2: Headers
    headers = [
        "CoCd",
        "G/L",
        "Year/month",
        "Type",
        "Reference",
        "DocumentNo",
        "Supplier",
        "Supplier Name",
        "PO Number",
        "Tax Based",
        "WHT",
        "Tax Rate",
        "URN",
        "User Name",
        "Text",
        "Clrng doc.",
        "Doc. Date",
        "Pstng Date",
        "Doc Curr",
        "Amount in doc. curr.",
        "Loc Curr",
        "Amount in local cur."
    ]
    
    # Style for headers
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Apply headers to row 2
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    ws.row_dimensions[2].height = 30
    
    # Row 3: Example data (optional, can be removed)
    example_data = [
        "1000",  # CoCd
        "150000",  # G/L
        "202601",  # Year/month
        "KR",  # Type
        "REF001",  # Reference
        "DOC001",  # DocumentNo
        "V001",  # Supplier
        "Sample Vendor Inc.",  # Supplier Name
        "PO-001",  # PO Number
        "10000000",  # Tax Based
        "200000",  # WHT
        "2",  # Tax Rate
        "URN001",  # URN
        "John Doe",  # User Name
        "Sample transaction",  # Text
        "CLR001",  # Clrng doc.
        "2026-01-15",  # Doc. Date
        "2026-01-15",  # Pstng Date
        "IDR",  # Doc Curr
        "10000000",  # Amount in doc. curr.
        "IDR",  # Loc Curr
        "10000000"  # Amount in local cur.
    ]
    
    # Apply example data to row 3 with light styling
    example_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for col_idx, value in enumerate(example_data, start=1):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = value
        cell.fill = example_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="left", vertical="center")
    
    # Row 4 onwards: Empty rows for data entry
    for row_idx in range(4, 8):
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
    
    # Set column widths
    column_widths = {
        'A': 8,   # CoCd
        'B': 10,  # G/L
        'C': 12,  # Year/month
        'D': 8,   # Type
        'E': 12,  # Reference
        'F': 12,  # DocumentNo
        'G': 10,  # Supplier
        'H': 20,  # Supplier Name
        'I': 12,  # PO Number
        'J': 15,  # Tax Based
        'K': 12,  # WHT
        'L': 10,  # Tax Rate
        'M': 12,  # URN
        'N': 15,  # User Name
        'O': 20,  # Text
        'P': 12,  # Clrng doc.
        'Q': 12,  # Doc. Date
        'R': 12,  # Pstng Date
        'S': 10,  # Doc Curr
        'T': 18,  # Amount in doc. curr.
        'U': 10,  # Loc Curr
        'V': 18   # Amount in local cur.
    }
    
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
    
    # Freeze panes at row 3 (keep headers visible)
    ws.freeze_panes = 'A3'
    
    # Save the template
    output_path = Path(__file__).parent.parent.parent / "client" / "public" / "templates" / "gl-template.xlsx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    
    print(f"✓ GL template generated successfully at: {output_path}")
    print(f"  Template includes {len(headers)} columns with example data")

if __name__ == "__main__":
    generate_gl_template()
